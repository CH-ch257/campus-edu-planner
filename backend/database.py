"""
文档管理模块：文档上传、文本提取、数据库存储
技术栈：SQLAlchemy + SQLite + PyPDF2 + python-docx
"""

import os
import uuid
import hashlib
from datetime import datetime
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, text
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker

# ========== 数据库连接 ==========
DATABASE_URL = "sqlite:///./campus.db"
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# 上传文件保存目录
UPLOAD_DIR = "./uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# ========== 文档表模型 ==========

class Document(Base):
    """文档表：存储上传的课件、作业等资料"""
    __tablename__ = "documents"

    id = Column(Integer, primary_key=True, index=True)
    filename = Column(String(255), nullable=False, comment="文件名")
    filetype = Column(String(20), nullable=False, comment="文件类型：pdf / word")
    file_path = Column(String(500), comment="文件保存路径")
    content = Column(Text, comment="提取的文本内容")
    uploaded_at = Column(DateTime, default=datetime.now, comment="上传时间")

    category = Column(String(50), default="未分类", comment="文档自动分类")
    category_confirmed = Column(Integer, default=0, comment="0=AI自动分类，1=管理员人工修正过")

    version = Column(Integer, default=1, comment="文档版本号")
    is_latest = Column(Integer, default=1, comment="是否最新版本：1=是，0=历史版本")
    parent_doc_id = Column(Integer, nullable=True, comment="关联上一版本的文档id")
    teacher_id = Column(Integer, nullable=True, comment="上传教师id（归属校验用）")

class Chunk(Base):
    """文档分块表，用于RAG向量检索"""
    __tablename__ = "chunks"

    id = Column(Integer, primary_key=True, index=True)
    doc_id = Column(Integer, nullable=False, comment="关联documents表id")
    chunk_text = Column(Text, comment="分块文本")
    embedding = Column(Text, comment="向量json字符串")

class QuestionAnswer(Base):
    """问答记录表：AI问答留痕"""
    __tablename__ = "qa_records"

    id = Column(Integer, primary_key=True, index=True)
    user_id = Column(Integer, nullable=False, index=True, comment="提问用户id")
    question = Column(Text, nullable=False, comment="用户问题")
    answer = Column(Text, nullable=False, comment="AI回答")
    sources = Column(Text, comment="引用来源json：[{doc_id, doc_name}]")
    source_doc_ids = Column(Text, comment="引用来源文档id，逗号分隔")
    note = Column(Text, default="", comment="备注")
    permission = Column(String(20), default="private", comment="权限：private/link/public")
    share_token = Column(String(32), nullable=True, index=True, comment="分享token")
    created_at = Column(DateTime, default=datetime.now, comment="提问时间")

class User(Base):
    """用户表：注册/登录账号"""
    __tablename__ = "users"

    id = Column(Integer, primary_key=True, index=True)
    username = Column(String(50), unique=True, nullable=False, index=True, comment="用户名")
    password_hash = Column(String(128), nullable=False, comment="加盐哈希后的密码")
    role = Column(String(20), default="student", comment="角色：student/teacher/admin")
    status = Column(Integer, default=1, comment="账号状态：1=启用 0=禁用")
    token = Column(String(64), nullable=True, comment="登录token")
    created_at = Column(DateTime, default=datetime.now, comment="注册时间")

# ========== 建表 ==========

def hash_password(password: str) -> str:
    """对密码加盐做SHA256哈希（与 main.py 保持一致）"""
    salt = uuid.uuid4().hex
    digest = hashlib.sha256((salt + password).encode("utf-8")).hexdigest()
    return f"{salt}:{digest}"


def init_db():
    """初始化数据库，创建表，并做轻量迁移与预设固定账号"""
    Base.metadata.create_all(bind=engine)
    # 轻量迁移：为阶段一已有的 users 表补充新列（create_all 不会改已有表）
    with engine.begin() as conn:
        cols = [r[1] for r in conn.execute(text("PRAGMA table_info(users)"))]
        if "role" not in cols:
            conn.execute(text("ALTER TABLE users ADD COLUMN role VARCHAR(20) DEFAULT 'student'"))
        if "status" not in cols:
            conn.execute(text("ALTER TABLE users ADD COLUMN status INTEGER DEFAULT 1"))
        if "token" not in cols:
            conn.execute(text("ALTER TABLE users ADD COLUMN token VARCHAR(64)"))
        dcols = [r[1] for r in conn.execute(text("PRAGMA table_info(documents)"))]
        if "teacher_id" not in dcols:
            conn.execute(text("ALTER TABLE documents ADD COLUMN teacher_id INTEGER"))
    # 预设固定账号：高级管理员 admin、老师 teacher（仅当不存在时创建，不开放注册）
    db = SessionLocal()
    try:
        if not db.query(User).filter(User.username == "admin").first():
            db.add(User(username="admin", password_hash=hash_password("admin123"), role="admin", status=1))
        if not db.query(User).filter(User.username == "teacher").first():
            db.add(User(username="teacher", password_hash=hash_password("teacher123"), role="teacher", status=1))
        db.commit()
    finally:
        db.close()
    print("✅ 数据库初始化完成")

def get_db():
    """获取数据库会话（FastAPI依赖注入）"""
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# ========== 文本提取函数 ==========

def extract_text_from_pdf(file_path: str) -> str:
    """从PDF文件提取文本"""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        return f"[PDF提取失败: {e}]"

def extract_text_from_word(file_path: str) -> str:
    """从Word文件提取文本"""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text.strip()
    except Exception as e:
        return f"[Word提取失败: {e}]"

# ========== 文档 CRUD ==========

def save_document(db, filename: str, filetype: str, file_path: str, content: str):
    """保存文档记录到数据库"""
    doc = Document(
        filename=filename,
        filetype=filetype,
        file_path=file_path,
        content=content
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc

def list_documents(db):
    """查询全部文档列表"""
    return db.query(Document).order_by(Document.uploaded_at.desc()).all()

def search_documents(db, keyword: str):
    """按文件名关键词搜索文档"""
    return db.query(Document).filter(
        Document.filename.contains(keyword)
    ).order_by(Document.uploaded_at.desc()).all()

def get_document_by_id(db, doc_id: int):
    """根据ID获取单篇文档"""
    return db.query(Document).filter(Document.id == doc_id).first()

def extract_text_from_txt(file_path: str) -> str:
    """从txt文件提取文本"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read().strip()
    except Exception as e:
        return f"[TXT提取失败: {e}]"
